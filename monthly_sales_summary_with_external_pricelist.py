import os
import glob
import pandas as pd
import tkinter as tk
from tkinter import messagebox
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

def process_data(year, month):
    folder_path = r"C:\Users\szabo\Desktop\pythonscripts\production"
    found_pricelist = glob.glob(os.path.join(folder_path, "pricelist*.xls*"))
    if not found_pricelist:
        messagebox.showerror("Error", "No pricelist found!")
        return
    pricelist_file = found_pricelist[0]

    pricelist = pd.read_excel(pricelist_file)
    pricelist["Date"] = pd.to_datetime(pricelist["Date"]).dt.date

    summarized_df = pd.DataFrame()
    files = glob.glob(os.path.join(folder_path, "*production*site*.xls*"))
    if not files:
        messagebox.showwarning("Warning", "No matching files found.")
        return

    for file in files:
        try:
            site = os.path.splitext(os.path.basename(file))[0].split("_site_")[-1]
            df = pd.read_excel(file, header=0)
            df = df.loc[:, ~df.columns.str.contains('^Unnamed')]
            if df.empty:
                continue

            # Detect date column
            date_column = None
            for col in df.columns:
                if pd.to_datetime(df[col], errors='coerce').notna().any():
                    date_column = col
                    df[col] = pd.to_datetime(df[col]).dt.date
                    break
            if date_column is None:
                continue

            # Detect quantity column
            quantity_column = None
            for col in df.columns:
                if "quantity" in str(col).lower():
                    quantity_column = col
                    break
            if quantity_column is None:
                messagebox.showerror("Error", f"{file} - 'Quantity' column not found.")
                return

            # Check Product column
            if "Product" not in df.columns:
                messagebox.showerror("Error", f"{file} - 'Product' column not found.")
                return

            # Filter by selected year and month
            df = df[(pd.to_datetime(df[date_column]).dt.year == year) &
                    (pd.to_datetime(df[date_column]).dt.month == month)].copy()
            if df.empty:
                continue

            # Assign price from pricelist
            def assign_price(date, product):
                valid = pricelist[pricelist["Date"] <= date]
                if valid.empty or product not in valid.columns:
                    return 0
                return valid[product].iloc[(valid["Date"] <= date).sum() - 1]

            df["Price"] = df.apply(lambda x: assign_price(x[date_column], x["Product"]), axis=1)
            df["Income"] = df[quantity_column] * df["Price"]
            df["Site"] = site
            column_order = [c for c in df.columns if c != "Site"] + ["Site"]
            df = df[column_order]

            summarized_df = pd.concat([summarized_df, df], ignore_index=True)

        except Exception as e:
            print(f"Error while processing file: {file}\n{e}")
            continue

    if summarized_df.empty:
        messagebox.showinfo("Result", "No data for the selected period.")
        return

    output_file_name = f"Monthly_Sales_Report_{month:02d}{year}.xlsx"
    output_file_path = os.path.join(folder_path, output_file_name)

    if os.path.exists(output_file_path):
        answer = messagebox.askyesno("Warning", f"The file already exists:\n{output_file_path}\nDo you want to overwrite it?")
        if not answer:
            return

    write_excel(summarized_df, output_file_path, date_column, quantity_column)
    messagebox.showinfo("Done", f"Report generated:\n{output_file_path}")


def write_excel(summarized_df, output_file_path, date_column, quantity_column):
    with pd.ExcelWriter(output_file_path, engine="openpyxl") as writer:
        # ===== Summary sheet =====
        summary_df = summarized_df.groupby("Site")[[quantity_column, "Income"]].sum().reset_index()
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

        # ===== Daily breakdown =====
        daily_df = summarized_df.groupby([date_column, "Site"])[[quantity_column, "Income"]].sum().reset_index()
        daily_df.to_excel(writer, sheet_name="Daily Breakdown", index=False)

        # ===== Site-specific sheets =====
        for site in summarized_df["Site"].unique():
            site_df = summarized_df[summarized_df["Site"] == site].copy()
            sheet_name = str(site)[:31]  # Excel sheet max length = 31
            site_df.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.book[sheet_name]

            # Add total row
            last_row = sheet.max_row + 1
            sheet[f"A{last_row}"] = "Total"
            for idx, cell in enumerate(sheet[1], start=1):
                col_letter = get_column_letter(idx)
                if cell.value == quantity_column or cell.value == "Income":
                    sheet[f"{col_letter}{last_row}"] = f"=SUM({col_letter}2:{col_letter}{last_row-1})"
                    if cell.value == "Income":
                        sheet[f"{col_letter}{last_row}"].number_format = '#,##0 "€"'

        workbook = writer.book

        # ===== Intelligent column formatting =====
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]

            # Center header row
            for cell in sheet[1]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Adjust column widths
            for col in sheet.columns:
                col_letter = get_column_letter(col[0].column)
                max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
                sheet.column_dimensions[col_letter].width = max_length + 5  # add padding

            # Format numbers and dates
            for idx, cell in enumerate(sheet[1], start=1):
                col_letter = get_column_letter(idx)
                if cell.value.lower() == "income":
                    for row in range(2, sheet.max_row + 1):
                        sheet[f"{col_letter}{row}"].number_format = '#,##0 "€"'
                elif cell.value.lower() in ["price", "quantity"]:
                    for row in range(2, sheet.max_row + 1):
                        sheet[f"{col_letter}{row}"].number_format = '#,##0'
                elif cell.value == date_column:
                    for row in range(2, sheet.max_row + 1):
                        sheet[f"{col_letter}{row}"].number_format = 'DD-MM-YYYY'


# ===== GUI =====
def start():
    try:
        monthyear = monthyear_entry.get().strip()
        if len(monthyear) != 6 or not monthyear.isdigit():
            raise ValueError("Please enter the period in MMYYYY format (e.g., 022026).")
        year = int(monthyear[2:])
        month = int(monthyear[:2])
        if month < 1 or month > 12:
            raise ValueError("Month should be between 1 and 12.")
        window.destroy()
        process_data(year, month)
    except Exception as e:
        messagebox.showerror("Error", str(e))


window = tk.Tk()
window.title("Monthly Sales Report Generator")

tk.Label(window, text="MonthYear (MMYYYY, e.g., 022026):").grid(row=0, column=0, padx=10, pady=5)
monthyear_entry = tk.Entry(window)
monthyear_entry.grid(row=0, column=1, padx=10, pady=5)
tk.Button(window, text="START", command=start).grid(row=1, columnspan=2, pady=10)

window.mainloop()
